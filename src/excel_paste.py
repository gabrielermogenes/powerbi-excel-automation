"""
Etapa 2 - Colagem dos arquivos exportados na planilha .xlsm.

Usa o próprio Excel (xlwings) para preservar macros, gráficos e formatação.
Antes de escrever, valida se as colunas do arquivo batem com a tabela de
destino. Se algo falhar no meio da escrita, a planilha NÃO é salva e o
arquivo original permanece intacto.
"""

import os
import shutil
import unicodedata
from datetime import datetime

import config


class ErroDeValidacao(Exception):
    """Problema achado ANTES de mexer na aba: nada foi alterado nela."""


def normalizar(texto):
    """Compara cabeçalhos ignorando maiúscula, espaço sobrando e acento
    (assim "Início" e "Inicio" contam como iguais)."""
    texto = " ".join(str(texto if texto is not None else "").split()).lower()
    return "".join(
        c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)
    )


def ler_exportacao(caminho):
    """
    Lê o arquivo exportado pelo Power BI e devolve (cabecalho, dados).

    A linha de cabeçalho é achada sozinha: é a que começa com o texto definido
    em PRIMEIRA_COLUNA_DO_CABECALHO (acima dela o Power BI coloca uma linha de
    "Filtros aplicados"). As colunas são as do cabeçalho, até a primeira
    célula vazia.
    """
    import openpyxl  # import aqui dentro: a etapa 1 funciona mesmo sem o openpyxl instalado

    # Lemos o arquivo inteiro (sem read_only): os arquivos do Power BI declaram
    # por dentro que só têm a célula A1, e no modo read_only o openpyxl
    # acredita nisso e enxerga uma célula só.
    wb = openpyxl.load_workbook(caminho, data_only=True)
    abas = wb.sheetnames
    linhas = [list(l) for l in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()

    indice_cab = None
    for i, linha in enumerate(linhas[:30]):
        if linha and normalizar(linha[0]).startswith(normalizar(config.PRIMEIRA_COLUNA_DO_CABECALHO)):
            indice_cab = i
            break
    if indice_cab is None:
        print(f"[DIAGNÓSTICO] Abas do arquivo: {abas} (li a primeira)")
        print(f"[DIAGNÓSTICO] Linhas lidas: {len(linhas)}")
        for i, linha in enumerate(linhas[:8], start=1):
            print(f"   linha {i}: {[('' if c is None else str(c)[:25]) for c in linha[:5]]}")
        raise ErroDeValidacao(
            f'Não achei a linha de cabeçalho (a que começa com "{config.PRIMEIRA_COLUNA_DO_CABECALHO}") '
            f"nas primeiras linhas de {os.path.basename(caminho)}."
        )

    cabecalho = []
    for celula in linhas[indice_cab]:
        if celula is None or str(celula).strip() == "":
            break
        cabecalho.append(str(celula).strip())
    equivalentes = {normalizar(k): v for k, v in config.CABECALHOS_EQUIVALENTES.items()}
    cabecalho = [equivalentes.get(normalizar(c), c) for c in cabecalho]
    n = len(cabecalho)

    dados = []
    for linha in linhas[indice_cab + 1:]:
        linha = (list(linha) + [None] * n)[:n]
        if all(v is None or str(v).strip() == "" for v in linha):
            break  # linha totalmente vazia: acabaram os dados
        if normalizar(linha[0]) == "total":
            continue  # linha de total não entra
        dados.append(linha)
    return cabecalho, dados


def escolher_tabela(candidatas, cabecalho_export):
    """
    candidatas: lista de (nome, [cabeçalhos da tabela]).
    Devolve a posição da tabela cujos primeiros N cabeçalhos são iguais aos do
    arquivo (N = nº de colunas do arquivo). Se nenhuma bater, levanta
    ErroDeValidacao mostrando as diferenças. Necessário porque uma aba pode ter
    várias tabelas.
    """
    if not candidatas:
        raise ErroDeValidacao("A aba não tem nenhuma tabela do Excel.")

    alvo = [normalizar(c) for c in cabecalho_export]
    for i, (nome, cabs) in enumerate(candidatas):
        if [normalizar(c) for c in cabs[:len(alvo)]] == alvo:
            return i

    msg = ["Nenhuma tabela da aba tem as mesmas colunas do arquivo. Nada foi colado:"]
    parecidas = [(nome, cabs) for nome, cabs in candidatas if cabs and normalizar(cabs[0]) == alvo[0]]
    for nome, cabs in (parecidas or candidatas)[:3]:
        difs = []
        for k in range(len(alvo)):
            d = cabs[k] if k < len(cabs) else "(sem coluna)"
            if k >= len(cabs) or normalizar(cabs[k]) != alvo[k]:
                difs.append(f"coluna {k + 1}: planilha '{d}' | arquivo '{cabecalho_export[k]}'")
        msg.append(f"   tabela '{nome}': " + "; ".join(difs[:6]))
    raise ErroDeValidacao("\n".join(msg))


def listar_tabelas(ws):
    """Devolve [(nome, [cabeçalhos])] de todas as tabelas da aba (via Excel)."""
    candidatas = []
    for i in range(1, ws.api.ListObjects.Count + 1):
        lo = ws.api.ListObjects(i)
        valores = ws.range(lo.HeaderRowRange.Address).value
        if not isinstance(valores, list):
            valores = [valores]
        candidatas.append((lo.Name, ["" if v is None else str(v) for v in valores]))
    return candidatas


def abrir_planilha(caminho):
    """
    Devolve (app, wb, abri_agora). Se a planilha já está aberta no Excel,
    usa a que está aberta (abri_agora = False) em vez de tentar abrir de novo.
    """
    import xlwings as xw  # import aqui dentro: só precisa do xlwings quem for usar a etapa 2

    caminho = os.path.abspath(caminho)
    for app in xw.apps:
        for livro in app.books:
            if os.path.normcase(livro.fullname) == os.path.normcase(caminho):
                return app, livro, False

    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    livro = app.books.open(caminho, update_links=False)
    return app, livro, True


def preparar_excel(app):
    """Cálculo manual e tela parada durante a colagem (bem mais rápido com
    milhares de linhas e fórmulas). Devolve o estado anterior pra restaurar."""
    estado = {}
    try:
        estado["calculo"] = app.api.Calculation
        estado["tela"] = app.api.ScreenUpdating
        app.api.Calculation = -4135  # xlCalculationManual
        app.api.ScreenUpdating = False
    except Exception:
        pass
    return estado


def restaurar_excel(app, estado):
    try:
        if "calculo" in estado:
            app.api.Calculation = estado["calculo"]
        if "tela" in estado:
            app.api.ScreenUpdating = estado["tela"]
    except Exception:
        pass


def colar_aba(wb, nome_aba, cabecalho_export, dados):
    nomes_abas = [s.name for s in wb.sheets]
    if nome_aba not in nomes_abas:
        raise ErroDeValidacao(f"A planilha não tem a aba '{nome_aba}'. Abas existentes: {nomes_abas}")
    ws = wb.sheets[nome_aba]

    # --- acha a tabela certa pelos cabeçalhos (e confere que as colunas batem) ---
    indice = escolher_tabela(listar_tabelas(ws), cabecalho_export)
    tabela = ws.api.ListObjects(indice + 1)

    if len(dados) == 0:
        raise ErroDeValidacao("O arquivo exportado não tem nenhuma linha de dados. Nada foi colado.")

    n = len(cabecalho_export)
    linha_cab = tabela.HeaderRowRange.Row
    col_ini = tabela.HeaderRowRange.Column
    ultima_col = col_ini + n - 1
    print(f"   Tabela '{tabela.Name}': colando {n} de {tabela.ListColumns.Count} colunas "
          "(as outras, de fórmula, ficam como estão).")

    try:
        if tabela.AutoFilter is not None and tabela.AutoFilter.FilterMode:
            print("   ATENÇÃO: essa tabela está com filtro ativo. Os dados entram em todas as "
                  "linhas, mas linhas que o filtro escondia podem continuar escondidas.")
    except Exception:
        pass

    # --- cresce a tabela se precisar ---
    capacidade = tabela.ListRows.Count
    if len(dados) > capacidade:
        total_colunas = tabela.ListColumns.Count
        nova_area = ws.range(
            (linha_cab, col_ini), (linha_cab + len(dados), col_ini + total_colunas - 1)
        )
        tabela.Resize(nova_area.api)
        capacidade = len(dados)
        print(f"   A tabela cresceu para {capacidade} linhas.")

    # --- limpa só os valores antigos das colunas coladas, sem mexer em formatação nem em fórmulas ---
    primeira_linha = linha_cab + 1
    ws.range((primeira_linha, col_ini), (primeira_linha + capacidade - 1, ultima_col)).clear_contents()

    # --- cola os dados novos ---
    ws.range((primeira_linha, col_ini)).value = dados
    print(f"   {len(dados)} linhas coladas (a partir da linha {primeira_linha}).")


def rodar_macros(app, wb, nomes):
    """
    Roda as macros da planilha, na ordem, depois de colar tudo. Se uma falhar,
    o erro é impresso e a próxima continua. Nunca derruba o script.
    """
    print("\n--- Rodando as macros da planilha ---")
    try:
        app.api.Calculate()  # garante que as fórmulas já refletem os dados colados
        # Se a macro abrir uma caixa de mensagem (MsgBox), com o Excel escondido
        # ela ficaria invisível e o script travaria esperando um OK. Por isso
        # deixamos o Excel visível durante as macros.
        app.visible = True
    except Exception as e:
        print(f"   Aviso: não consegui preparar o Excel pras macros: {e}")

    for nome in nomes:
        try:
            print(f"   Rodando '{nome}'... (se aparecer uma caixa de mensagem, clique em OK)")
            # Chamamos direto pelo Excel: aceita tanto macro de módulo ("NomeDaMacro")
            # quanto macro de aba ("Planilha1.NomeDaMacro")
            app.api.Run(f"'{wb.name}'!{nome}")
            # Se a macro escreve o resultado na barra de status do Excel, mostramos aqui
            status = app.api.StatusBar
            if isinstance(status, str) and status:
                print(f"   {status}")
            app.api.StatusBar = False  # devolve a barra de status ao normal
            print("   OK.")
        except Exception as e:
            print(f"   ERRO na macro '{nome}': {e}")
            print("   Se a macro não foi encontrada: confira o nome e se o módulo VBA está dentro da planilha.")
            print("   Se a mensagem fala de macros desabilitadas/bloqueadas: desbloqueie o arquivo")
            print("   (botão direito no .xlsm > Propriedades > Desbloquear) ou use uma pasta confiável.")


def colar_na_planilha(exportados):
    """
    Etapa 2. `exportados` é a lista de arquivos que a etapa 1 exportou COM
    SUCESSO nesta rodada, ou None se a etapa 1 foi desligada (aí usa o que já
    está na pasta). Cada aba é independente: se uma falhar ANTES de ser
    alterada, as outras continuam. Se alguma falhar no meio da escrita, a
    planilha NÃO é salva (o arquivo original fica intacto).
    """
    print("\n=== ETAPA 2: COLANDO NA PLANILHA ===")
    try:
        if "COLE_AQUI" in config.PLANILHA_DESTINO:
            raise ErroDeValidacao("Preencha PLANILHA_DESTINO no config.py com o caminho da planilha .xlsm.")
        if not os.path.exists(config.PLANILHA_DESTINO):
            raise ErroDeValidacao(f"Não achei a planilha: {config.PLANILHA_DESTINO}")

        # quais arquivos entram nesta rodada
        a_colar = []
        for arquivo, aba in config.ABAS_PARA_COLAR:
            if exportados is not None and arquivo not in exportados:
                print(f"   {arquivo}: a extração falhou nesta rodada, então a aba '{aba}' NÃO será "
                      "atualizada (o arquivo antigo da pasta não é confiável).")
            elif not os.path.exists(os.path.join(config.PASTA_EXPORTACOES, arquivo)):
                print(f"   {arquivo}: não encontrado na pasta, aba '{aba}' não será atualizada.")
            else:
                a_colar.append((arquivo, aba))
        if not a_colar:
            raise ErroDeValidacao("Nenhum arquivo disponível pra colar.")

        if config.FAZER_BACKUP:
            base, ext = os.path.splitext(config.PLANILHA_DESTINO)
            backup = f"{base}_backup_{datetime.now():%Y%m%d_%H%M%S}{ext}"
            shutil.copy2(config.PLANILHA_DESTINO, backup)
            print(f"Backup salvo em: {backup}")

        print("Abrindo a planilha no Excel...")
        app, wb, abri_agora = abrir_planilha(config.PLANILHA_DESTINO)
        estado_excel = preparar_excel(app)
        resultados = []
        falha_na_escrita = False
        try:
            for arquivo, aba in a_colar:
                print(f"\n--- {arquivo} -> aba '{aba}' ---")
                try:
                    caminho_export = os.path.join(config.PASTA_EXPORTACOES, arquivo)
                    gerado_em = datetime.fromtimestamp(os.path.getmtime(caminho_export))
                    print(f"   Arquivo gerado em {gerado_em:%d/%m/%Y %H:%M}")
                    if gerado_em.date() != datetime.now().date():
                        print("   ATENÇÃO: esse arquivo NÃO é de hoje.")
                    try:
                        cabecalho, dados = ler_exportacao(caminho_export)
                    except ErroDeValidacao:
                        raise
                    except Exception as e:
                        raise ErroDeValidacao(f"não consegui ler o arquivo: {e}")
                    print(f"   {len(dados)} linhas de dados, {len(cabecalho)} colunas.")
                    colar_aba(wb, aba, cabecalho, dados)
                    resultados.append((aba, "OK"))
                except ErroDeValidacao as e:
                    print(f"   PULADA (nada foi alterado nessa aba): {e}")
                    resultados.append((aba, "pulada"))
                except Exception as e:
                    print(f"   ERRO ao escrever: {e}")
                    falha_na_escrita = True
                    resultados.append((aba, "ERRO"))
        except Exception:
            falha_na_escrita = True
            raise
        finally:
            restaurar_excel(app, estado_excel)
            algum_ok = any(status == "OK" for _, status in resultados)
            if config.MACROS_DEPOIS_DE_COLAR and algum_ok and not falha_na_escrita:
                rodar_macros(app, wb, config.MACROS_DEPOIS_DE_COLAR)
            if abri_agora:
                try:
                    if falha_na_escrita or not algum_ok:
                        wb.api.Close(SaveChanges=False)
                        print("\nPlanilha fechada SEM salvar" +
                              (" (houve erro durante a escrita; o arquivo original continua intacto)."
                               if falha_na_escrita else " (nada foi colado)."))
                    else:
                        wb.save()
                        wb.close()
                        print("\nPlanilha salva.")
                finally:
                    app.quit()
            else:
                print("\nA planilha já estava aberta no Excel, então NÃO salvei "
                      "(pra não salvar nada que você estivesse editando). Confira e salve você.")

        print("\nRESUMO:")
        for aba, status in resultados:
            print(f"   {aba}: {status}")
    except Exception as e:
        print(f"\nERRO na etapa 2: {e}")
